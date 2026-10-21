# SPDX-License-Identifier: Apache-2.0

"""Small helpers shared by Desk families: CSV and XLSX text, dates, rounding."""

import csv
import datetime as dt
import io
from decimal import ROUND_HALF_UP, Decimal


def csv_text(rows: list[dict] | list[list], header: list[str]) -> str:
    """CSV text with a header line. ``rows`` are dicts keyed by header or plain lists."""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(header)
    for row in rows:
        writer.writerow([row[h] for h in header] if isinstance(row, dict) else row)
    return buf.getvalue()


def read_csv(text: str | bytes) -> list[dict[str, str]]:
    """Rows of CSV text as dicts (a leading BOM is ignored)."""
    if isinstance(text, bytes):
        text = text.decode("utf-8")
    return list(csv.DictReader(io.StringIO(text.lstrip("﻿"))))


def xlsx_bytes(sheets: dict[str, list[list]]) -> bytes:
    """An .xlsx workbook with one sheet per item; the first row of each sheet is its header."""
    from openpyxl import Workbook

    book = Workbook()
    first = True
    for name, rows in sheets.items():
        sheet = book.active if first else book.create_sheet()
        assert sheet is not None
        sheet.title = name
        first = False
        for row in rows:
            sheet.append(row)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def read_xlsx(data: bytes) -> dict[str, list[list]]:
    """Every sheet of an .xlsx workbook as rows of cell values."""
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    return {sheet.title: [list(r) for r in sheet.iter_rows(values_only=True)] for sheet in book}


def half_up(value: float | Decimal, places: int = 2) -> float:
    """Round half up (0.125 -> 0.13 at 2 places), computed in Decimal from the shortest repr."""
    quantum = Decimal(1).scaleb(-places)
    return float(
        Decimal(repr(value) if isinstance(value, float) else value).quantize(
            quantum, rounding=ROUND_HALF_UP
        )
    )


def day(text: str) -> dt.date:
    return dt.date.fromisoformat(text)


def iso(date: dt.date) -> str:
    return date.isoformat()
