# SPDX-License-Identifier: Apache-2.0

"""
Task contract and output grading for Workbench Desk.

A family builds a ``DeskTask``: input files, an instruction that names every
input and output path, and the outputs the grader checks. Outputs are files the
agent writes into its workspace: a JSON object, a CSV table or one plain value.

Grading reads each output file (a regular file of at most MAX_OUTPUT_BYTES, no
symlinks). An exact answer earns 1.0. Otherwise, the format and correct-value
fractions give partial credit capped at ``PARTIAL_REWARD_CAP``.
"""

import csv
import io
import json
import math
import os
import random
import re
import stat
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import dataclass, field
from typing import Any

FORMAT_CREDIT = 0.1
PARTIAL_VALUE_BONUS = 0.1
PARTIAL_REWARD_CAP = 0.2
MAX_OUTPUT_BYTES = 256 * 1024
CHECK_KINDS = frozenset({"int", "num", "text", "set", "list", "map", "bool"})


@dataclass(frozen=True)
class Check:
    """How one value is compared. ``tol`` and ``rel`` apply to numbers, in lists and maps too."""

    kind: str
    tol: float = 0.0
    rel: float = 0.0

    def __post_init__(self) -> None:
        if self.kind not in CHECK_KINDS:
            raise ValueError(f"unknown check kind: {self.kind}")


def _is_num(value: Any) -> bool:
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        return False
    try:
        return math.isfinite(value)
    except OverflowError:
        return False


def _num_ok(check: Check, got: Any, want: Any) -> bool:
    if not _is_num(got):
        return False
    limit = max(check.tol, check.rel * abs(float(want)))
    return abs(float(got) - float(want)) <= limit + 1e-9


def _scalar_ok(check: Check, got: Any, want: Any) -> bool:
    if isinstance(want, bool):
        return got is want
    if _is_num(want):
        return _num_ok(check, got, want)
    return isinstance(got, str) and got.strip() == str(want).strip()


def _scalar_key(value: Any) -> str:
    if _is_num(value) and float(value).is_integer():
        return json.dumps(int(value))
    if isinstance(value, str):
        return json.dumps(value.strip())
    return json.dumps(value)


def value_ok(check: Check, got: Any, want: Any) -> bool:
    """True when ``got`` equals ``want`` under ``check``."""
    kind = check.kind
    if kind == "int":
        return _is_num(got) and float(got).is_integer() and int(got) == int(want)
    if kind == "num":
        return _num_ok(check, got, want)
    if kind == "text":
        return isinstance(got, str) and got.strip() == str(want).strip()
    if kind == "bool":
        return isinstance(got, bool) and got is bool(want)
    if kind == "set":
        if not isinstance(got, list) or len(got) != len(want):
            return False
        return sorted(map(_scalar_key, got)) == sorted(map(_scalar_key, want))
    if kind == "list":
        if not isinstance(got, list) or len(got) != len(want):
            return False
        return all(_scalar_ok(check, g, w) for g, w in zip(got, want, strict=True))
    if kind == "map":
        if not isinstance(got, dict):
            return False
        got_norm = {str(k).strip(): v for k, v in got.items()}
        if set(got_norm) != {str(k) for k in want}:
            return False
        return all(_scalar_ok(check, got_norm[str(k)], w) for k, w in want.items())
    raise ValueError(kind)


def shape_ok(check: Check, got: Any) -> bool:
    """True when ``got`` has the type ``check`` asks for, whatever its value."""
    kind = check.kind
    if kind in ("int", "num"):
        return _is_num(got)
    if kind == "text":
        return isinstance(got, str)
    if kind == "bool":
        return isinstance(got, bool)
    if kind in ("set", "list"):
        return isinstance(got, list)
    return isinstance(got, dict)


def _expected_units(check: Check, want: Any) -> int:
    """How many independently checkable values the expected answer contains."""
    if check.kind in ("list", "set", "map"):
        return max(len(want), 1)
    return 1


def _credit_units(check: Check, got: Any, want: Any) -> tuple[int, int, int]:
    """Return (well-shaped, correct, total) value units for a JSON field."""
    if check.kind in ("list", "set"):
        total = (
            max(len(got), len(want), 1) if isinstance(got, list) else _expected_units(check, want)
        )
        if not isinstance(got, list):
            return 0, 0, total
        if not got and not want:
            return 1, 1, 1
        if check.kind == "set":
            actual = Counter(map(_scalar_key, got))
            expected = Counter(map(_scalar_key, want))
            correct = sum((actual & expected).values())
        else:
            correct = sum(
                _scalar_ok(check, got[i], value) for i, value in enumerate(want) if i < len(got)
            )
        return total, correct, total

    if check.kind == "map":
        if not isinstance(got, dict):
            return 0, 0, _expected_units(check, want)
        got_norm = {str(key).strip(): value for key, value in got.items()}
        total = max(len(got_norm), len(want), 1)
        if not got_norm and not want:
            return 1, 1, 1
        correct = sum(
            key in got_norm and _scalar_ok(check, got_norm[key], value)
            for key, value in want.items()
        )
        return total, correct, total

    total = 1
    shaped = int(shape_ok(check, got))
    correct = int(value_ok(check, got, want))
    return shaped, correct, total


def _parse_cell(check: Check, cell: str) -> Any:
    """A CSV cell as the type the column asks for, or None when it does not parse."""
    text = cell.strip()
    if check.kind in ("int", "num"):
        try:
            value = float(text)
        except ValueError:
            return None
        return value if math.isfinite(value) else None
    return text


@dataclass(frozen=True)
class JsonOut:
    """A JSON object file. Extra keys are allowed; every key in ``checks`` must be right."""

    path: str
    checks: dict[str, Check]
    want: dict[str, Any]

    def grade(self, text: str) -> tuple[bool, bool]:
        try:
            got = json.loads(text)
        except (ValueError, RecursionError):
            return False, False
        if not isinstance(got, dict):
            return False, False
        shape = all(name in got and shape_ok(c, got[name]) for name, c in self.checks.items())
        right = shape and all(
            value_ok(c, got[name], self.want[name]) for name, c in self.checks.items()
        )
        return shape, right

    def score_units(self, text: str) -> tuple[int, int, int]:
        """Count shape and exact-value credit without discarding correct fields."""
        expected_total = sum(
            _expected_units(check, self.want[name]) for name, check in self.checks.items()
        )
        try:
            got = json.loads(text)
        except (ValueError, RecursionError):
            return 0, 0, expected_total
        if not isinstance(got, dict):
            return 0, 0, expected_total
        shaped = correct = total = 0
        for name, check in self.checks.items():
            shape_units, correct_units, field_units = _credit_units(
                check, got.get(name), self.want[name]
            )
            shaped += shape_units
            correct += correct_units
            total += field_units
        return shaped, correct, total


@dataclass(frozen=True)
class CsvOut:
    """A CSV file with an exact header. Rows match in order, or by ``key`` column when set."""

    path: str
    header: tuple[str, ...]
    columns: tuple[Check, ...]
    rows: tuple[tuple[Any, ...], ...]
    key: int | None = None

    def grade(self, text: str) -> tuple[bool, bool]:
        try:
            lines = [
                row for row in csv.reader(io.StringIO(text.lstrip("\ufeff"), newline="")) if row
            ]
        except csv.Error:
            return False, False
        if not lines or tuple(c.strip() for c in lines[0]) != self.header:
            return False, False
        body = lines[1:]
        if len(body) != len(self.rows) or any(len(r) != len(self.header) for r in body):
            return False, False
        parsed = [
            tuple(_parse_cell(c, x) for c, x in zip(self.columns, r, strict=True)) for r in body
        ]
        if any(v is None for r in parsed for v in r):
            return False, False
        if self.key is None:
            pairs = list(zip(parsed, self.rows, strict=True))
        else:
            k = self.key
            by_key = {str(r[k]).strip(): r for r in parsed}
            if len(by_key) != len(parsed):
                return True, False
            pairs = []
            for want in self.rows:
                got = by_key.get(str(want[k]).strip())
                if got is None:
                    return True, False
                pairs.append((got, want))
        right = all(
            all(_scalar_ok(c, g, w) for c, g, w in zip(self.columns, got, want, strict=True))
            for got, want in pairs
        )
        return True, right

    def score_units(self, text: str) -> tuple[int, int, int]:
        """Score parseable cells individually so one wrong cell does not erase the rest."""
        expected_cells = max(len(self.rows) * len(self.header), 1)
        try:
            lines = [
                row for row in csv.reader(io.StringIO(text.lstrip("\ufeff"), newline="")) if row
            ]
        except csv.Error:
            return 0, 0, expected_cells
        if not lines or tuple(cell.strip() for cell in lines[0]) != self.header:
            return 0, 0, expected_cells

        body = lines[1:]
        width = max(len(self.header), max((len(row) for row in body), default=0), 1)
        total = max(len(self.rows), len(body), 1) * width
        shaped = correct = 0

        if self.key is None:
            for row_index in range(max(len(self.rows), len(body))):
                got_row = body[row_index] if row_index < len(body) else []
                want_row = self.rows[row_index] if row_index < len(self.rows) else None
                for col_index in range(width):
                    if col_index >= len(got_row):
                        continue
                    check = self.columns[col_index] if col_index < len(self.columns) else None
                    if check is None:
                        continue
                    value = _parse_cell(check, got_row[col_index])
                    if value is None:
                        continue
                    shaped += 1
                    if want_row is not None and col_index < len(want_row):
                        correct += int(_scalar_ok(check, value, want_row[col_index]))
            return shaped, correct, total

        key = self.key
        by_key: dict[str, list[str]] = {}
        for row in body:
            if key < len(row):
                by_key.setdefault(row[key].strip(), row)

        # A valid CSV still earns format credit for cells whose row key is unexpected.
        for row in body:
            for col_index, raw in enumerate(row[: len(self.columns)]):
                if _parse_cell(self.columns[col_index], raw) is not None:
                    shaped += 1

        for want_row in self.rows:
            got_row = by_key.get(str(want_row[key]).strip())
            if got_row is None:
                continue
            for col_index, check in enumerate(self.columns):
                if col_index >= len(got_row):
                    continue
                value = _parse_cell(check, got_row[col_index])
                if value is not None:
                    correct += int(_scalar_ok(check, value, want_row[col_index]))
        return min(shaped, total), min(correct, total), total


@dataclass(frozen=True)
class TextOut:
    """A file holding one value: a number, an integer or a line of text."""

    path: str
    check: Check
    want: Any

    def grade(self, text: str) -> tuple[bool, bool]:
        value = _parse_cell(self.check, text)
        if value is None or (self.check.kind == "text" and not value):
            return False, False
        return True, value_ok(self.check, value, self.want)

    def score_units(self, text: str) -> tuple[int, int, int]:
        value = _parse_cell(self.check, text)
        if value is None or (self.check.kind == "text" and not value):
            return 0, 0, 1
        return 1, int(value_ok(self.check, value, self.want)), 1


@dataclass(frozen=True)
class XlsxOut:
    """An .xlsx workbook. Each ``cells`` key is "Sheet!B12"; the cell must hold a value
    (a number or text, not a formula without a cached result) that ``check`` accepts."""

    path: str
    cells: dict[str, tuple[Check, Any]]

    def grade(self, data: bytes) -> tuple[bool, bool]:
        from openpyxl import load_workbook

        try:
            book = load_workbook(io.BytesIO(data), data_only=True)
        except Exception:
            return False, False
        shape, right = True, True
        for ref, (check, want) in self.cells.items():
            sheet, cell = ref.split("!")
            if sheet not in book.sheetnames:
                return False, False
            value = book[sheet][cell].value
            if value is None or not shape_ok(check, value):
                shape, right = False, False
            elif not value_ok(check, value, want):
                right = False
        return shape, right

    def score_units(self, data: bytes) -> tuple[int, int, int]:
        from openpyxl import load_workbook

        total = max(len(self.cells), 1)
        try:
            book = load_workbook(io.BytesIO(data), data_only=True)
        except Exception:
            return 0, 0, total
        shaped = correct = 0
        for ref, (check, want) in self.cells.items():
            sheet, cell = ref.split("!")
            if sheet not in book.sheetnames:
                continue
            value = book[sheet][cell].value
            if not shape_ok(check, value):
                continue
            shaped += 1
            correct += int(value_ok(check, value, want))
        return shaped, correct, total


Output = JsonOut | CsvOut | TextOut | XlsxOut


@dataclass
class DeskTask:
    """One task instance. ``instruction`` writes the workspace directory as ``{root}``."""

    instruction: str
    files: dict[str, str | bytes]
    outputs: list[Output]
    oracle: dict[str, str | bytes]
    decoys: list[dict[str, str | bytes]] = field(default_factory=list)

    def text(self, root: str) -> str:
        return self.instruction.replace("{root}", root)


class DeskFamily:
    """One task type. ``build`` makes an instance; ``solve`` re-derives its outputs from files."""

    id: str = ""
    domain: str = ""
    title: str = ""
    levels: tuple[int, ...] = (1, 2)

    def build(self, level: int, rng: random.Random) -> DeskTask:
        raise NotImplementedError

    def solve(
        self, files: Mapping[str, str | bytes], instruction: str
    ) -> Mapping[str, str | bytes]:
        raise NotImplementedError


def read_output(workdir: str, relpath: str) -> bytes | None:
    """The bytes of an output file, or None when it is missing, too big or not a regular file."""
    path = os.path.join(workdir, relpath)
    try:
        info = os.lstat(path)
    except OSError:
        return None
    if not stat.S_ISREG(info.st_mode) or info.st_size > MAX_OUTPUT_BYTES:
        return None
    try:
        fd = os.open(path, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK)
        with os.fdopen(fd, "rb") as fh:
            return fh.read(MAX_OUTPUT_BYTES)
    except OSError:
        return None


def grade_texts(
    outputs: Sequence[Output], texts: Mapping[str, str | bytes | None]
) -> tuple[float, str]:
    """Reward for output file contents keyed by path, with a short note for the agent."""
    shaped = correct = total = 0
    for out in outputs:
        content = texts.get(out.path)
        if content is None:
            return 0.0, f"{out.path} is missing"
        if isinstance(out, XlsxOut):
            data = content if isinstance(content, bytes) else content.encode("utf-8")
            shape_units, correct_units, total_units = out.score_units(data)
        else:
            if isinstance(content, bytes):
                content = content.decode("utf-8", errors="replace")
            if isinstance(out, JsonOut):
                shape_units, correct_units, total_units = out.score_units(content)
            elif isinstance(out, CsvOut):
                shape_units, correct_units, total_units = out.score_units(content)
            else:
                shape_units, correct_units, total_units = out.score_units(content)
        shaped += shape_units
        correct += correct_units
        total += total_units
    if total == 0:
        return 0.0, "there are no checked output values"
    if correct == total:
        return 1.0, "all outputs are correct"
    reward = min(
        PARTIAL_REWARD_CAP,
        round((FORMAT_CREDIT * shaped + PARTIAL_VALUE_BONUS * correct) / total, 8),
    )
    if shaped == total and correct == 0:
        return FORMAT_CREDIT, "the outputs have the right shape but some values are wrong"
    if correct:
        return reward, "some output values are correct"
    if shaped:
        return reward, "some output values have the expected format"
    return reward, "the outputs do not have the asked format"


def grade_workdir(outputs: Sequence[Output], workdir: str) -> tuple[float, str]:
    return grade_texts(outputs, {o.path: read_output(workdir, o.path) for o in outputs})


def must_match(pattern: str, text: str, flags: int = 0) -> re.Match[str]:
    """Search ``text`` for ``pattern``. A missing match is a bug in the family, so raise."""
    match = re.search(pattern, text, flags)
    if match is None:
        raise ValueError(f"no match for {pattern!r}")
    return match


def money(value: float) -> float:
    """Round half up to the cent, as accountants do (0.005 -> 0.01)."""
    from decimal import ROUND_HALF_UP, Decimal

    return float(Decimal(repr(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
